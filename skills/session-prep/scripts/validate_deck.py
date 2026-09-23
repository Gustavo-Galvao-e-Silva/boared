# /// script
# requires-python = ">=3.10"
# dependencies = ["panchi>=2.0.0", "python-pptx>=1.0", "pyyaml>=6.0", "pydantic>=2.0", "sympy>=1.12", "pillow>=10"]
# ///
"""Check a built deck before it is delivered. Exit 1 on any problem.

Usage:
    validate_deck.py deck.pptx [--deck deck.yaml] [--template template.pptx --map template-map.yaml]

Checks:
  - no [bracket] placeholder text is left (slides, tables and notes)
  - every picture's image is present in the file
  - no media file is orphaned (in the package but referenced by nothing)
  with --deck:
  - one slide per deck.yaml entry, so no template-original slides remain
  - every slide's hidden flag matches the spec (and, with --template, the template)
  - when deck.yaml sits next to questions.yaml: verify.log passed and is current,
    so nothing ships unverified
"""

from __future__ import annotations

import argparse
import posixpath
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from build_deck import expected_hidden, is_hidden, resolve_source  # noqa: E402
from pptx_common import iter_shapes  # noqa: E402
from schemas import SchemaError, TemplateMap, load_deck, load_template_map  # noqa: E402

# [Insert statement], [Type the question here] — but not [1, 2] or [A | b]
PLACEHOLDER = re.compile(r"\[(?=[^\]\n]*[A-Za-z]{3})[^\]\n]{3,}\]")


def _texts(slide):
    for shape in iter_shapes(slide.shapes):
        if shape.has_text_frame:
            yield shape.name, shape.text_frame.text
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield f"{shape.name} (table)", cell.text
    if slide.has_notes_slide:
        yield "notes", slide.notes_slide.notes_text_frame.text


def orphan_media(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        media = {n for n in names if n.startswith("ppt/media/")}
        referenced = set()
        for rels in (n for n in names if n.endswith(".rels")):
            base = posixpath.dirname(posixpath.dirname(rels))  # ppt/slides/_rels/x.rels → ppt/slides
            for target in (t.decode() for t in re.findall(rb'Target="([^"]+)"', z.read(rels))):
                path = target[1:] if target.startswith("/") else posixpath.join(base, target)
                referenced.add(posixpath.normpath(path))
    return sorted(media - referenced)


def validate(pptx: Path, deck_yaml: Path | None = None, template: Path | None = None, map_path: Path | None = None) -> list[str]:
    problems = []
    prs = Presentation(str(pptx))
    slides = list(prs.slides)

    for n, slide in enumerate(slides, start=1):
        for where, text in _texts(slide):
            for m in PLACEHOLDER.finditer(text):
                problems.append(f"slide {n}: placeholder text {m.group(0)!r} left in {where}")
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    _ = shape.image.blob
                except Exception as e:
                    problems.append(f"slide {n}: picture {shape.name!r} has no image data ({e})")
    problems += [f"orphan media file {m}" for m in orphan_media(pptx)]

    if deck_yaml is not None:
        try:
            deck = load_deck(deck_yaml)
            tmap = load_template_map(map_path) if map_path else TemplateMap()
        except SchemaError as e:
            return problems + [str(e)]
        if len(slides) != len(deck.slides):
            problems.append(
                f"{len(slides)} slides but deck.yaml has {len(deck.slides)} — template-original slides left in, or slides missing"
            )
        template_slides = list(Presentation(str(template)).slides) if template else None
        for n, (slide, spec) in enumerate(zip(slides, deck.slides), start=1):
            if template_slides is not None:
                number, kind = resolve_source(spec, tmap)
                want = expected_hidden(spec, kind, template_slides[number - 1])
            elif spec.hidden is not None:
                want = spec.hidden
            else:
                continue
            if is_hidden(slide) != want:
                problems.append(f"slide {n}: hidden is {is_hidden(slide)}, spec says {want}")

        session = deck_yaml.resolve().parent
        if (session / "questions.yaml").exists():
            from verify_runner import log_problem

            problem = log_problem(session)
            if problem:
                problems.append(f"not verified: {problem}")
    return problems


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--deck", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--map", type=Path, dest="map_path")
    args = parser.parse_args()
    problems = validate(args.pptx, args.deck, args.template, args.map_path)
    for p in problems:
        print(f"FAIL {p}")
    print(f"{args.pptx.name}: {'OK' if not problems else f'{len(problems)} problem(s)'}")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
