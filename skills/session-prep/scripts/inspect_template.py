"""Describe a .pptx template so its slides can be mapped to session slide kinds.

Usage:
    inspect_template.py template.pptx                 # human-readable dump
    inspect_template.py template.pptx --draft-map OUT # also write a draft template-map.yaml

The dump lists, for every slide: its 1-based number, layout, and every shape
(name, id, type, position in inches, and a preview of its text). The draft map
has one entry per slide with every text shape listed as a field; the skill then
renames entries to kinds (warmup, question, tf, ...) and trims fields with the
user's confirmation.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Emu
from pptx_common import iter_shapes


def _inches(v: int | None) -> float | None:
    return None if v is None else round(Emu(v).inches, 2)


def describe(path: Path) -> list[dict]:
    prs = Presentation(str(path))
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        shapes = []
        for s in iter_shapes(slide.shapes):
            text = s.text_frame.text.strip() if s.has_text_frame else ""
            shapes.append(
                {
                    "name": s.name,
                    "id": s.shape_id,
                    "type": str(s.shape_type).split(".")[-1].split(" ")[0] if s.shape_type else "UNKNOWN",
                    "box_in": [_inches(s.left), _inches(s.top), _inches(s.width), _inches(s.height)],
                    "has_text": s.has_text_frame,
                    "text": text,
                }
            )
        slides.append({"number": number, "layout": slide.slide_layout.name, "shapes": shapes})
    return slides


def print_dump(path: Path, slides: list[dict]) -> None:
    prs = Presentation(str(path))
    print(f"{path}  —  {len(slides)} slides, {_inches(prs.slide_width)}x{_inches(prs.slide_height)} in\n")
    for sl in slides:
        print(f"Slide {sl['number']}  [layout: {sl['layout']}]")
        for s in sl["shapes"]:
            preview = s["text"].replace("\n", " / ")
            if len(preview) > 90:
                preview = preview[:87] + "..."
            box = ", ".join("?" if v is None else str(v) for v in s["box_in"])
            print(f"  - {s['name']!r} (id {s['id']}, {s['type']}, box {box})" + (f": {preview}" if preview else ""))
        print()


def draft_map(template: Path, slides: list[dict]) -> dict:
    kinds = {}
    for sl in slides:
        fields = {}
        for s in sl["shapes"]:
            if s["has_text"]:
                key = s["name"].lower().replace(" ", "_")
                fields[key] = s["id"]
        kinds[f"slide{sl['number']}"] = {
            "slide": sl["number"],
            "description": (sl["shapes"][0]["text"][:60] if sl["shapes"] else ""),
            "fields": fields,
            "math_area": None,
        }
    return {"template": template.name, "kinds": kinds}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("template", type=Path)
    parser.add_argument("--draft-map", type=Path)
    args = parser.parse_args()

    slides = describe(args.template)
    print_dump(args.template, slides)
    if args.draft_map:
        args.draft_map.write_text(yaml.safe_dump(draft_map(args.template, slides), sort_keys=False, allow_unicode=True))
        print(f"Draft map written to {args.draft_map}")


if __name__ == "__main__":
    main()
