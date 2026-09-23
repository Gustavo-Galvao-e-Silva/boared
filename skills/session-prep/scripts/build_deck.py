"""Build a session deck by copying template slides and replacing their text.

Never designs slides: every output slide is a copy of a template slide, with
text swapped into existing shapes (keeping their fonts, colors, bullets) and
math images fitted into a designated area.

Usage:
    build_deck.py deck.yaml --template template.pptx --map template-map.yaml -o deck.pptx
    build_deck.py deck.yaml -o deck.pptx          # no template: plain 16:9 fallback

deck.yaml:
    slides:
      - kind: question            # entry in template-map.yaml `kinds`
        # or: from_slide: 4       # copy template slide 4 directly (fields = shape names/ids)
        text:
          title: "Possible or Impossible?"
          body: |                 # one paragraph per line; formatting cloned from the template
            A 3x3 matrix with two equal columns that is invertible.
            A 2x2 matrix A with A^2 = 0 and A != 0.
        images:
          - path: math/rref.png   # relative to deck.yaml
            area: math_area       # map key, shape name, or shape id whose box the image fills
            # or box: [left, top, width, height]   (inches)
        notes: "Answers: Impossible (det = 0); Possible, e.g. [[0,1],[0,0]]"

template-map.yaml (see inspect_template.py --draft-map):
    kinds:
      question:
        slide: 4                  # 1-based slide number in the template
        fields: {title: "Title 1", body: 7}   # field -> shape name or id
        math_area: "Rectangle 9"  # optional default image area (removed from output)
"""

from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches
from pptx_common import find_shape

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class DeckError(Exception):
    pass


# --- slide copying -----------------------------------------------------------


def duplicate_slide(prs, source):
    """Append a copy of ``source`` (same layout, shapes, background, images)."""
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):  # drop layout placeholders; we copy the source's own
        shape._element.getparent().remove(shape._element)

    rid_map = {}
    for rel in source.part.rels.values():
        if rel.reltype in (RT.NOTES_SLIDE, RT.SLIDE_LAYOUT):
            continue
        if rel.is_external:
            rid_map[rel.rId] = dest.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = dest.part.relate_to(rel.target_part, rel.reltype)

    src_tree = source.shapes._spTree
    dest_tree = dest.shapes._spTree
    for el in src_tree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        dest_tree.append(copy.deepcopy(el))

    src_bg = source._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        dest._element.cSld.insert(0, copy.deepcopy(src_bg))

    for el in dest._element.iter():
        for attr, value in list(el.attrib.items()):
            if attr.startswith(f"{{{_R_NS}}}") and value in rid_map:
                el.set(attr, rid_map[value])
    return dest


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


# --- content -----------------------------------------------------------------


def set_text(shape, value: str) -> None:
    """Replace text, one paragraph per line, cloning the template paragraphs' formatting."""
    if not shape.has_text_frame:
        raise DeckError(f"Shape {shape.name!r} has no text frame")
    txBody = shape.text_frame._txBody
    originals = txBody.findall(qn("a:p"))
    lines = str(value).rstrip("\n").split("\n")

    for i, line in enumerate(lines):
        p = copy.deepcopy(originals[min(i, len(originals) - 1)])
        runs = p.findall(qn("a:r"))
        for extra in p.findall(qn("a:fld")) + p.findall(qn("a:br")) + runs[1:]:
            p.remove(extra)
        if runs:
            runs[0].find(qn("a:t")).text = line
        else:  # empty template paragraph: build a run using its end-paragraph formatting
            r = p.makeelement(qn("a:r"), {})
            end = p.find(qn("a:endParaRPr"))
            if end is not None:
                rpr = copy.deepcopy(end)
                rpr.tag = qn("a:rPr")
                r.append(rpr)
            t = r.makeelement(qn("a:t"), {})
            t.text = line
            r.append(t)
            end_idx = list(p).index(end) if end is not None else len(p)
            p.insert(end_idx, r)
        txBody.append(p)

    for p in originals:
        txBody.remove(p)


def place_image(slide, image: Path, box: tuple[int, int, int, int]) -> None:
    """Add ``image`` fitted (aspect preserved) and centered inside ``box`` (EMU)."""
    left, top, width, height = box
    pic = slide.shapes.add_picture(str(image), left, top)
    scale = min(width / pic.width, height / pic.height)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = left + (width - pic.width) // 2
    pic.top = top + (height - pic.height) // 2


def default_box(prs) -> tuple[int, int, int, int]:
    """Lower 60% of the slide, with margins — used when no area is given."""
    w, h = prs.slide_width, prs.slide_height
    return (int(w * 0.08), int(h * 0.35), int(w * 0.84), int(h * 0.58))


# --- build -------------------------------------------------------------------


def resolve_source(spec: dict, template_map: dict) -> tuple[int, dict, object]:
    """Return (1-based template slide number, field->shape map, default math area) for a deck slide."""
    if "from_slide" in spec:
        return int(spec["from_slide"]), {}, spec.get("math_area")
    kind = spec.get("kind")
    kinds = template_map.get("kinds", {})
    if kind not in kinds:
        raise DeckError(f"Unknown kind {kind!r}. Known kinds: {', '.join(kinds) or '(none — no template map)'}")
    entry = kinds[kind]
    return int(entry["slide"]), entry.get("fields") or {}, entry.get("math_area")


def build(deck_path: Path, out: Path, template: Path | None, map_path: Path | None) -> Path:
    deck = yaml.safe_load(deck_path.read_text())
    base = deck_path.parent
    slides_spec = deck.get("slides") or []
    if not slides_spec:
        raise DeckError("deck.yaml has no slides")

    if template is None:
        return build_plain(slides_spec, base, out)

    template_map = yaml.safe_load(map_path.read_text()) if map_path else {}
    prs = Presentation(str(template))
    originals = list(prs.slides)

    for n, spec in enumerate(slides_spec, start=1):
        where = f"deck slide {n}"
        try:
            number, fields, math_area = resolve_source(spec, template_map)
            if not 1 <= number <= len(originals):
                raise DeckError(f"template slide {number} does not exist (template has {len(originals)})")
            slide = duplicate_slide(prs, originals[number - 1])

            for field, value in (spec.get("text") or {}).items():
                set_text(find_shape(slide, fields.get(field, field)), value)

            area_shapes = {}
            for img in spec.get("images") or []:
                path = (base / img["path"]).resolve()
                if not path.exists():
                    raise DeckError(f"image not found: {path}")
                if "box" in img:
                    box = tuple(int(Inches(v)) for v in img["box"])
                else:
                    area_ref = img.get("area", math_area)
                    if area_ref is None:
                        box = default_box(prs)
                    else:
                        area = find_shape(slide, fields.get(area_ref, area_ref) if isinstance(area_ref, str) else area_ref)
                        box = (area.left, area.top, area.width, area.height)
                        area_shapes[area.shape_id] = area
                place_image(slide, path, box)
            for area in area_shapes.values():  # the area is a placeholder rectangle, not content
                if not (area.has_text_frame and area.text_frame.text.strip()):
                    area._element.getparent().remove(area._element)

            if spec.get("notes"):
                slide.notes_slide.notes_text_frame.text = str(spec["notes"])
        except (DeckError, KeyError) as e:
            raise DeckError(f"{where}: {e}") from e

    for slide in originals:
        delete_slide(prs, slide)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def build_plain(slides_spec: list[dict], base: Path, out: Path) -> Path:
    """Fallback when there is no template: title + body + images on a blank 16:9 deck."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[1]  # Title and Content
    for spec in slides_spec:
        slide = prs.slides.add_slide(layout)
        text = spec.get("text") or {}
        slide.shapes.title.text = str(text.get("title", ""))
        body = slide.placeholders[1]
        images = spec.get("images") or []
        body_text = "\n".join(str(v) for k, v in text.items() if k != "title")
        body.left, body.top, body.width = Inches(0.8), Inches(1.6), Inches(11.7)
        body.height = Inches(1.8) if images else Inches(5.4)
        if body_text:
            set_text(body, body_text)
        else:
            body._element.getparent().remove(body._element)
        for img in images:
            box = tuple(int(Inches(v)) for v in img["box"]) if "box" in img else default_box(prs)
            place_image(slide, (base / img["path"]).resolve(), box)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("deck", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--map", type=Path, dest="map_path")
    parser.add_argument("-o", "--out", type=Path, required=True)
    args = parser.parse_args()
    if args.map_path and not args.template:
        parser.error("--map requires --template")
    try:
        print(build(args.deck, args.out, args.template, args.map_path))
    except DeckError as e:
        sys.exit(f"build_deck: {e}")


if __name__ == "__main__":
    main()
