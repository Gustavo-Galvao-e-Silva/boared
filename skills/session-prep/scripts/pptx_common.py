"""Shared python-pptx helpers for inspect_template.py and build_deck.py."""

from __future__ import annotations

from collections.abc import Iterator

from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape


def iter_shapes(shapes) -> Iterator[BaseShape]:
    """Yield every shape on a slide, descending into groups."""
    for shape in shapes:
        yield shape
        if isinstance(shape, GroupShape):
            yield from iter_shapes(shape.shapes)


def find_shape(slide, ref: str | int) -> BaseShape:
    """Find a shape by name (str) or shape id (int). Raises KeyError if absent or ambiguous."""
    matches = [
        s
        for s in iter_shapes(slide.shapes)
        if (s.shape_id == ref if isinstance(ref, int) else s.name == ref)
    ]
    if not matches:
        available = ", ".join(f"{s.name!r} (id {s.shape_id})" for s in iter_shapes(slide.shapes))
        raise KeyError(f"No shape {ref!r} on slide. Available: {available}")
    if len(matches) > 1:
        ids = ", ".join(str(s.shape_id) for s in matches)
        raise KeyError(f"Shape name {ref!r} is ambiguous (ids {ids}); reference it by id instead.")
    return matches[0]
