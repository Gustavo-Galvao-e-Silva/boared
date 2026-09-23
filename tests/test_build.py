"""Golden tests for build_deck.py: fixture template → build → assertions on the XML."""

import pytest
import yaml
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import build_deck
import render_math
from conftest import BG, BRAND, TEMPLATE_MAP, needs_tex


def build(tmp_path, template, slides, tmap=TEMPLATE_MAP, warnings=None):
    (tmp_path / "template-map.yaml").write_text(yaml.safe_dump(tmap))
    (tmp_path / "deck.yaml").write_text(yaml.safe_dump({"slides": slides}, allow_unicode=True))
    out = build_deck.build(tmp_path / "deck.yaml", tmp_path / "deck.pptx", template, tmp_path / "template-map.yaml", warnings)
    return Presentation(str(out))


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def fill_of(sh):
    sppr = sh._element.find(qn("p:spPr"))
    solid = sppr.find(qn("a:solidFill"))
    if solid is not None:
        return solid.find(qn("a:srgbClr")).get("val")
    return "none" if sppr.find(qn("a:noFill")) is not None else "style"


def test_text_keeps_template_formatting(tmp_path, template):
    prs = build(tmp_path, template, [{"kind": "warmup", "text": {"title": "Brain dump", "body": "a\nb\nc"}}])
    (s1,) = prs.slides
    run = shape(s1, "Heading").text_frame.paragraphs[0].runs[0]
    assert run.text == "Brain dump"
    assert run.font.color.rgb == BRAND and run.font.bold and run.font.size == Pt(40)
    body = shape(s1, "Body")
    assert [p.text for p in body.text_frame.paragraphs] == ["a", "b", "c"]
    assert all(p.runs[0].font.size == Pt(24) for p in body.text_frame.paragraphs)
    assert s1.background.fill.fore_color.rgb == BG


def test_unset_fields_keep_template_text(tmp_path, template):
    prs = build(tmp_path, template, [{"kind": "warmup", "text": {"title": "Closing proof"}}])
    assert shape(prs.slides[0], "Body").text_frame.text == "First prompt\nSecond prompt"


def test_hidden_slides_stay_hidden(tmp_path, template):
    prs = build(
        tmp_path,
        template,
        [
            {"kind": "facilitator"},  # hidden in the template → hidden in the deck
            {"kind": "warmup"},
            {"kind": "warmup", "hidden": True},  # spec overrides
            {"kind": "tf-answers", "hidden": False},
        ],
    )
    assert [build_deck.is_hidden(s) for s in prs.slides] == [True, False, True, False]
    assert len(prs.slides) == 4  # template slides removed


def test_rich_text_markup(tmp_path, template):
    prs = build(
        tmp_path,
        template,
        [{"kind": "warmup", "text": {"body": "**Claim:** A is {{orange|invertible}}\nplain line\n{{00AA00|**both**}} end"}}],
    )
    paras = shape(prs.slides[0], "Body").text_frame.paragraphs
    runs = paras[0].runs
    assert [r.text for r in runs] == ["Claim:", " A is ", "invertible"]
    assert runs[0].font.bold and not runs[1].font.bold
    assert str(runs[2].font.color.rgb) == "E87722"
    assert all(r.font.size == Pt(24) for r in runs)  # template run formatting cloned per segment
    assert paras[1].runs[0].text == "plain line"
    both = paras[2].runs[0]
    assert both.text == "both" and both.font.bold and str(both.font.color.rgb) == "00AA00"


def test_unknown_colour_is_an_error(tmp_path, template):
    with pytest.raises(build_deck.DeckError, match="deck slide 1: unknown colour 'teal'"):
        build(tmp_path, template, [{"kind": "warmup", "text": {"body": "{{teal|x}}"}}])


def test_parse_markup():
    assert build_deck.parse_markup("a **b** {{red|c}} d") == [
        ("a ", False, None),
        ("b", True, None),
        (" ", False, None),
        ("c", False, "red"),
        (" d", False, None),
    ]
    assert build_deck.parse_markup("no markup") == [("no markup", False, None)]


def test_table_rows_use_first_data_row_font(tmp_path, template):
    rows = [
        ["1", "If A is invertible and A = LU, then A⁻¹ = L⁻¹U⁻¹."],
        ["2", "**Every** matrix has an LU factorization."],
        ["3", "x"],
    ]
    prs = build(tmp_path, template, [{"kind": "tf", "tables": {"statements": {"shape": "Statements", "rows": rows}}}])
    table = shape(prs.slides[0], "Statements").table
    assert len(table.rows) == 4
    assert table.cell(0, 1).text == "Statement"
    assert table.cell(0, 1).text_frame.paragraphs[0].runs[0].font.size == Pt(28)  # header kept
    for r in (1, 2, 3):
        run = table.cell(r, 1).text_frame.paragraphs[0].runs[0]
        assert run.font.size == Pt(20)  # data rows copy the first data row, not the header
    assert table.cell(2, 1).text_frame.paragraphs[0].runs[0].font.bold
    assert table.cell(3, 1).text == "x"


def test_table_width_mismatch(tmp_path, template):
    with pytest.raises(build_deck.DeckError, match="row 1 has 3 cells; the table has 2 columns"):
        build(tmp_path, template, [{"kind": "tf", "tables": {"statements": {"shape": "Statements", "rows": [["1", "a", "b"]]}}}])


def test_answer_marks_from_exemplars(tmp_path, template):
    prs = build(tmp_path, template, [{"kind": "tf-answers", "states": {"mark": ["off", "on"]}}])
    s = prs.slides[0]
    assert fill_of(shape(s, "Oval 1")) == "none"  # the 'on' exemplar itself switched off
    assert fill_of(shape(s, "Oval 2")) == "E87722"  # started with no fill, now filled
    assert "Oval 3" not in [sh.name for sh in s.shapes]  # unused mark removed


def test_choice_highlight(tmp_path, template):
    prs = build(tmp_path, template, [{"kind": "mcq-answer", "states": {"choice": {"on": ["B", "D"]}}}])
    s = prs.slides[0]
    assert [fill_of(shape(s, f"Choice {c}")) for c in "ABCD"] == ["none", "E87722", "none", "E87722"]
    b = shape(s, "Choice B").text_frame.paragraphs[0].runs[0]
    assert b.text == "B" and b.font.bold and str(b.font.color.rgb) == "E87722"
    a = shape(s, "Choice A").text_frame.paragraphs[0].runs[0]
    assert not a.font.bold


def test_unknown_state_and_cell(tmp_path, template):
    with pytest.raises(build_deck.DeckError, match="unknown state 'glow'"):
        build(tmp_path, template, [{"kind": "tf-answers", "states": {"glow": ["on"]}}])
    with pytest.raises(build_deck.DeckError, match="unknown cells"):
        build(tmp_path, template, [{"kind": "mcq-answer", "states": {"choice": {"on": ["E"]}}}])


def test_geometry_overrides(tmp_path, template):
    spec = {"text": "one\ntwo", "box": [1, 2, 6, 3], "anchor": "ctr", "font_size": 30, "line_spacing": 1.5}
    prs = build(tmp_path, template, [{"kind": "warmup", "text": {"body": spec}}])
    body = shape(prs.slides[0], "Body")
    assert (body.left, body.top, body.width, body.height) == (Inches(1), Inches(2), Inches(6), Inches(3))
    assert body.text_frame._txBody.find(qn("a:bodyPr")).get("anchor") == "ctr"
    assert all(p.runs[0].font.size == Pt(30) for p in body.text_frame.paragraphs)
    assert body.text_frame.paragraphs[0].line_spacing == 1.5


def test_fit_warning(tmp_path, template):
    warnings = []
    long = "\n".join(["A very long statement about invertible matrices and LU factorizations"] * 8)
    build(tmp_path, template, [{"kind": "warmup", "text": {"body": long}}], warnings=warnings)
    assert any("may overflow" in w and "Body" in w for w in warnings)
    warnings.clear()
    build(tmp_path, template, [{"kind": "warmup", "text": {"body": "short"}}], warnings=warnings)
    assert warnings == []


def test_images_fill_math_area(tmp_path, template):
    from PIL import Image

    (tmp_path / "math").mkdir()
    Image.new("RGBA", (400, 100), (0, 0, 0, 255)).save(tmp_path / "math" / "inv.png")
    prs = build(tmp_path, template, [{"kind": "warmup", "images": [{"path": "math/inv.png"}], "notes": "Yes: det = 1"}])
    s = prs.slides[0]
    assert "Math Area" not in [sh.name for sh in s.shapes]  # placeholder area replaced by the image
    assert sum(1 for sh in s.shapes if sh.shape_type == 13) == 2  # logo + math image
    assert s.notes_slide.notes_text_frame.text == "Yes: det = 1"


def test_build_errors_name_the_slide(tmp_path, template):
    (tmp_path / "deck.yaml").write_text(yaml.safe_dump({"slides": [{"from_slide": 2, "text": {"Nope": "x"}}]}))
    with pytest.raises(build_deck.DeckError, match="deck slide 1: .*No shape 'Nope'"):
        build_deck.build(tmp_path / "deck.yaml", tmp_path / "deck.pptx", template, None)


def test_schema_errors_are_reported(tmp_path, template):
    (tmp_path / "deck.yaml").write_text(yaml.safe_dump({"slides": [{"kind": "tf", "from_slide": 2}]}))
    with pytest.raises(build_deck.DeckError, match="either kind or from_slide"):
        build_deck.build(tmp_path / "deck.yaml", tmp_path / "deck.pptx", template, None)


def test_build_plain_fallback(tmp_path):
    (tmp_path / "deck.yaml").write_text(
        yaml.safe_dump({"slides": [{"text": {"title": "**T**", "body": "x\ny"}}, {"text": {"title": "H"}, "hidden": True}]})
    )
    prs = Presentation(str(build_deck.build(tmp_path / "deck.yaml", tmp_path / "deck.pptx", None, None)))
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title.text == "T"
    assert build_deck.is_hidden(prs.slides[1])


@needs_tex
def test_render_panchi_rref(tmp_path):
    import panchi as pan
    from panchi.algorithms import rref

    A = pan.exact_matrix([[1, 2, 3], [2, 5, 7], [0, 1, 2]])
    out = render_math.render(rref(A)._repr_latex_(), tmp_path / "rref.png")
    assert out.stat().st_size > 1000


@needs_tex
def test_render_reports_latex_errors(tmp_path):
    with pytest.raises(RuntimeError, match="LaTeX failed"):
        render_math.render(r"\begin{bmatrix} 1 \end{bmatri}", tmp_path / "bad.png")


def test_one_based_rows():
    assert render_math.one_based_rows(r"R_{0} \to R_{0} + (-2)\,R_{2}") == r"R_{1} \to R_{1} + (-2)\,R_{3}"


def test_strip_delimiters():
    assert render_math.strip_delimiters("$$x$$") == "x"
    assert render_math.strip_delimiters("$x$") == "x"
    assert render_math.strip_delimiters("x") == "x"
