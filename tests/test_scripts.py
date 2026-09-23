import shutil
import subprocess
import sys
from pathlib import Path

import panchi as pan
import pytest
import yaml
from panchi.algorithms import rref
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "session-prep" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import build_deck  # noqa: E402
import inspect_template  # noqa: E402
import render_math  # noqa: E402

needs_tex = pytest.mark.skipif(shutil.which("latex") is None, reason="no TeX install")
BRAND = RGBColor(0x7C, 0x3A, 0xED)


def make_template(path: Path, logo: Path) -> None:
    """Two-slide template: a warm-up slide and a question slide with a math area and a logo."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for heading in ("Warm-up", "Main session"):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFA, 0xF5, 0xFF)
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
        title.name = "Heading"
        run = title.text_frame.paragraphs[0].add_run()
        run.text = heading
        run.font.size, run.font.bold, run.font.color.rgb = Pt(40), True, BRAND
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12), Inches(1.5))
        body.name = "Body"
        for i, line in enumerate(["First prompt", "Second prompt"]):
            p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size = Pt(24)
        slide.shapes.add_picture(str(logo), Inches(12.3), Inches(6.6), Inches(0.8))
        area = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3.2), Inches(11), Inches(4))
        area.name = "Math Area"
    prs.save(str(path))


@pytest.fixture
def workspace(tmp_path):
    logo = tmp_path / "logo.png"
    if shutil.which("latex") is None:
        pytest.skip("no TeX install")
    render_math.render(r"\heartsuit", logo)
    template = tmp_path / "template.pptx"
    make_template(template, logo)
    return tmp_path, template


def test_one_based_rows():
    assert render_math.one_based_rows(r"R_{0} \to R_{0} + (-2)\,R_{2}") == r"R_{1} \to R_{1} + (-2)\,R_{3}"


def test_strip_delimiters():
    assert render_math.strip_delimiters("$$x$$") == "x"
    assert render_math.strip_delimiters("$x$") == "x"
    assert render_math.strip_delimiters("x") == "x"


@needs_tex
def test_render_panchi_rref(tmp_path):
    A = pan.exact_matrix([[1, 2, 3], [2, 5, 7], [0, 1, 2]])
    out = render_math.render(rref(A)._repr_latex_(), tmp_path / "rref.png")
    assert out.stat().st_size > 1000


@needs_tex
def test_render_reports_latex_errors(tmp_path):
    with pytest.raises(RuntimeError, match="LaTeX failed"):
        render_math.render(r"\begin{bmatrix} 1 \end{bmatri}", tmp_path / "bad.png")


def test_inspect_and_build(workspace):
    tmp, template = workspace
    slides = inspect_template.describe(template)
    assert [s["number"] for s in slides] == [1, 2]
    draft = inspect_template.draft_map(template, slides)
    assert "heading" in draft["kinds"]["slide1"]["fields"]

    (tmp / "template-map.yaml").write_text(
        yaml.safe_dump(
            {
                "kinds": {
                    "warmup": {"slide": 1, "fields": {"title": "Heading", "body": "Body"}},
                    "question": {
                        "slide": 2,
                        "fields": {"title": "Heading", "body": "Body"},
                        "math_area": "Math Area",
                    },
                }
            }
        )
    )
    A = pan.exact_matrix([[2, 1], [5, 3]])
    render_math.render(pan.inverse(A)._repr_latex_(), tmp / "math" / "inv.png")
    (tmp / "deck.yaml").write_text(
        yaml.safe_dump(
            {
                "slides": [
                    {"kind": "warmup", "text": {"title": "Brain dump", "body": "a\nb\nc"}},
                    {
                        "kind": "question",
                        "text": {"title": "Find the inverse", "body": "Is this right?"},
                        "images": [{"path": "math/inv.png"}],
                        "notes": "Yes: det = 1",
                    },
                    {"kind": "warmup", "text": {"title": "Closing proof"}},
                ]
            }
        )
    )
    out = build_deck.build(tmp / "deck.yaml", tmp / "deck.pptx", template, tmp / "template-map.yaml")

    prs = Presentation(str(out))
    assert len(prs.slides) == 3  # template slides removed
    s1, s2, s3 = prs.slides
    heading = next(s for s in s1.shapes if s.name == "Heading")
    run = heading.text_frame.paragraphs[0].runs[0]
    assert run.text == "Brain dump"
    assert run.font.color.rgb == BRAND and run.font.bold and run.font.size == Pt(40)
    body = next(s for s in s1.shapes if s.name == "Body")
    assert [p.text for p in body.text_frame.paragraphs] == ["a", "b", "c"]
    assert all(p.runs[0].font.size == Pt(24) for p in body.text_frame.paragraphs)
    assert s1.background.fill.fore_color.rgb == RGBColor(0xFA, 0xF5, 0xFF)

    names = [s.name for s in s2.shapes]
    assert "Math Area" not in names  # placeholder area replaced by the image
    assert sum(1 for s in s2.shapes if s.shape_type == 13) == 2  # logo + math image
    assert s2.notes_slide.notes_text_frame.text == "Yes: det = 1"
    assert next(s for s in s3.shapes if s.name == "Body").text_frame.text == "First prompt\nSecond prompt"


def test_build_errors_name_the_slide(workspace):
    tmp, template = workspace
    (tmp / "deck.yaml").write_text(yaml.safe_dump({"slides": [{"from_slide": 1, "text": {"Nope": "x"}}]}))
    with pytest.raises(build_deck.DeckError, match="deck slide 1: .*No shape 'Nope'"):
        build_deck.build(tmp / "deck.yaml", tmp / "deck.pptx", template, None)


def test_build_plain_fallback(tmp_path):
    (tmp_path / "deck.yaml").write_text(yaml.safe_dump({"slides": [{"text": {"title": "T", "body": "x\ny"}}]}))
    out = build_deck.build(tmp_path / "deck.yaml", tmp_path / "deck.pptx", None, None)
    assert len(Presentation(str(out)).slides) == 1


def test_verify_runner(tmp_path):
    good = tmp_path / "verify_good.py"
    good.write_text(
        "import panchi as pan\n"
        "from verify_kit import check\n"
        "A = pan.exact_matrix([[2, 1], [5, 3]])\n"
        "check('Q1 inverse', pan.inverse(A).inverse, pan.exact_matrix([[3, -1], [-5, 2]]))\n"
    )
    bad = tmp_path / "verify_bad.py"
    bad.write_text(good.read_text().replace("[[3, -1], [-5, 2]]", "[[3, 1], [-5, 2]]"))
    runner = SCRIPTS / "verify_runner.py"
    ok = subprocess.run([sys.executable, runner, good], capture_output=True, text=True)
    assert ok.returncode == 0, ok.stdout + ok.stderr
    assert "1 passed" in ok.stdout
    fail = subprocess.run([sys.executable, runner, bad], capture_output=True, text=True)
    assert fail.returncode == 1
    assert "FAIL Q1 inverse" in fail.stdout


def test_diff_decks(tmp_path):
    import diff_decks

    def deck(name, slides):
        (tmp_path / f"{name}.yaml").write_text(yaml.safe_dump({"slides": slides}))
        return build_deck.build(tmp_path / f"{name}.yaml", tmp_path / f"{name}.pptx", None, None)

    warm = {"text": {"title": "Warm-up", "body": "Brain dump: invertible matrices"}}
    tf = {"text": {"title": "True or False", "body": "AB = AC implies B = C\nA^2 = 0 implies A = 0"}}
    close = {"text": {"title": "Closing", "body": "Prove the IMT direction (a) => (d)"}}
    new_slide = {"text": {"title": "Quick poll", "body": "Which step was hardest?"}}
    tf_edited = {"text": {"title": "True or False", "body": "AB = AC implies B = C\nIf A is invertible, so is A^T"}}

    generated = deck("generated", [warm, tf, close])
    assert "No changes" in diff_decks.diff(generated, generated)

    final = deck("final", [warm, new_slide, tf_edited])
    report = diff_decks.diff(generated, final)
    assert "## Added slide 2: Quick poll" in report
    assert "## Slide 2 → 3: True or False" in report
    assert "-A^2 = 0 implies A = 0" in report and "+If A is invertible, so is A^T" in report
    assert "## Removed slide 3: Closing" in report
    assert "Warm-up" not in report.split("presented slides.")[1]  # unchanged slide not reported
