import shutil
import sys
from pathlib import Path

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "session-prep" / "scripts"
sys.path.insert(0, str(SCRIPTS))

BRAND = RGBColor(0x7C, 0x3A, 0xED)
ORANGE = RGBColor(0xE8, 0x77, 0x22)
BG = RGBColor(0xFA, 0xF5, 0xFF)

needs_tex = pytest.mark.skipif(shutil.which("latex") is None, reason="no TeX install")


def _heading(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    box.name = "Heading"
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb = Pt(40), True, BRAND
    return box


def _body(slide, lines=("First prompt", "Second prompt"), top=1.4, height=1.5):
    body = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12), Inches(height))
    body.name = "Body"
    for i, line in enumerate(lines):
        p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(24)
    return body


def _oval(slide, name, left, top, filled, text=""):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(0.5), Inches(0.5))
    o.name = name
    if filled:
        o.fill.solid()
        o.fill.fore_color.rgb = ORANGE
    else:
        o.fill.background()
    if text:
        r = o.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(18)
    return o


def make_template(path: Path, logo: Path) -> None:
    """Fixture template:
    1 title · 2 warm-up (math area, logo) · 3 T/F with a table · 4 T/F answers (hidden, marks)
    5 MCQ answer (choice ovals) · 6 hidden facilitator slide
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        return s

    s1 = slide()
    _heading(s1, "Session [N]")
    sub = _body(s1, ["[Topics]"])
    sub.name = "Subtitle"

    s2 = slide()
    _heading(s2, "Warm-up")
    _body(s2)
    s2.shapes.add_picture(str(logo), Inches(12.3), Inches(6.6), Inches(0.8))
    area = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3.2), Inches(11), Inches(4))
    area.name = "Math Area"
    s2.notes_slide.notes_text_frame.text = "Warm-up: brain dump"

    for number, hidden in ((3, False), (4, True)):
        s = slide()
        _heading(s, "True or False?" if number == 3 else "True or False? answers")
        frame = s.shapes.add_table(3, 2, Inches(0.5), Inches(1.5), Inches(10), Inches(2.4))
        frame.name = "Statements"
        table = frame.table
        for c, text in enumerate(["#", "Statement"]):
            cell = table.cell(0, c)
            cell.text = text
            cell.text_frame.paragraphs[0].runs[0].font.bold = True
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        for r in (1, 2):
            for c, text in enumerate([str(r), f"[Statement {r}]"]):
                cell = table.cell(r, c)
                cell.text = text
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        if hidden:
            _oval(s, "Oval 1", 11, 2.0, True)
            _oval(s, "Oval 2", 11, 2.8, False)
            _oval(s, "Oval 3", 11, 3.6, False)
            s._element.set("show", "0")

    s5 = slide()
    _heading(s5, "Warm-up answer")
    for i, letter in enumerate("ABCD"):
        o = _oval(s5, f"Choice {letter}", 1 + 2 * i, 3, letter == "A", letter)
        if letter == "A":
            o.text_frame.paragraphs[0].runs[0].font.bold = True

    s6 = slide()
    _heading(s6, "Facilitator notes")
    s6._element.set("show", "0")
    prs.save(str(path))


TEMPLATE_MAP = {
    "colors": {"orange": "E87722"},
    "kinds": {
        "title": {"slide": 1, "fields": {"title": "Heading", "subtitle": "Subtitle"}},
        "warmup": {"slide": 2, "fields": {"title": "Heading", "body": "Body"}, "math_area": "Math Area"},
        "tf": {"slide": 3, "fields": {"title": "Heading"}, "tables": {"statements": "Statements"}},
        "tf-answers": {
            "slide": 4,
            "fields": {"title": "Heading"},
            "tables": {"statements": "Statements"},
            "states": {"mark": {"cells": ["Oval 1", "Oval 2", "Oval 3"], "on": "Oval 1", "off": "Oval 2"}},
        },
        "mcq-answer": {
            "slide": 5,
            "fields": {"title": "Heading"},
            "states": {
                "choice": {
                    "cells": {"A": "Choice A", "B": "Choice B", "C": "Choice C", "D": "Choice D"},
                    "on": {"shape": "Choice A", "text_color": "orange", "bold": True},
                    "off": "Choice B",
                }
            },
        },
        "facilitator": {"slide": 6, "fields": {"title": "Heading"}},
    },
    "slots": {
        "warmup": {"kind": "warmup", "per_slide": 1},
        "tf": {
            "kind": "tf",
            "per_slide": 3,
            "table": "statements",
            "field": None,
            "answers": "tf-answers",
            "answer_state": "mark",
        },
        "pi": {"kind": "warmup", "per_slide": 2, "title": "Possible or Impossible?"},
        "problem": {"kind": "warmup", "per_slide": 1, "title": "Problem {n}"},
        "proof": {"kind": "warmup", "per_slide": 1, "title": "Closing"},
    },
}


@pytest.fixture
def template(tmp_path) -> Path:
    logo = tmp_path / "logo.png"
    Image.new("RGBA", (64, 64), (124, 58, 237, 255)).save(logo)
    path = tmp_path / "template.pptx"
    make_template(path, logo)
    return path


@pytest.fixture
def course(tmp_path, template) -> Path:
    """A course folder with the fixture template, its map, and a schedule with session numbers."""
    root = tmp_path / "course"
    (root / "sessions").mkdir(parents=True)
    shutil.copy(template, root / "template.pptx")
    (root / "template-map.yaml").write_text(yaml.safe_dump(TEMPLATE_MAP, sort_keys=False))
    (root / "course.yaml").write_text(
        yaml.safe_dump(
            {
                "course": "MATH 1554 — Linear Algebra",
                "term": "Fall 2026",
                "session": {"days": ["Tue", "Thu"], "length_min": 50},
                "schedule": [
                    {
                        "week": 5,
                        "start": "2026-09-21",
                        "sections": ["2.5 Matrix factorizations"],
                        "goals": ["compute an LU factorization"],
                        "sessions": [{"date": "2026-09-22", "number": 9}, {"date": "2026-09-24", "number": 10}],
                    }
                ],
            },
            sort_keys=False,
        )
    )
    (root / "lessons.md").write_text("# Lessons — MATH 0000, Fall 2026\n\n## Preferences\n")
    return root
